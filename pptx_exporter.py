

import asyncio
import math
import os
import re
import shutil
import subprocess
import tempfile
import time
from datetime import datetime, timedelta, timezone
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import parse_qs, urlparse

import edge_tts
import gspread
import requests
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pydrive2.drive import GoogleDrive

from services.sheets_service import get_sheet_data_and_df
from services.smart_progress_bar import SmartProgressBar

# Try to apply nest_asyncio so Edge TTS works inside already-running event loops (Streamlit)
try:
    import nest_asyncio
    nest_asyncio.apply()
except ImportError:
    pass


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

DEFAULT_WORKSHEET = "Slide Chunks"
FINAL_GRAPHICS_COLUMN = "final_graphics_definition"
DEFAULT_VOICE = "en-CA-LiamNeural"
SLIDE_WIDTH_INCHES = 13.333
SLIDE_HEIGHT_INCHES = 7.5
MAX_WORKERS = 8

# Presentation dimensions in EMU (English Metric Units)
SLIDE_WIDTH_EMU = Inches(SLIDE_WIDTH_INCHES)
SLIDE_HEIGHT_EMU = Inches(SLIDE_HEIGHT_INCHES)


def _io_worker_count() -> int:
    """
    Pick a practical worker count for IO-heavy tasks.
    Keeps enough parallelism for downloads/TTS metadata without oversubscription.
    """
    cpu = os.cpu_count() or 4
    return max(4, min(16, cpu * 2))


# ---------------------------------------------------------------------------
# Text / URL helpers
# ---------------------------------------------------------------------------

def _safe_str(value) -> str:
    if value is None:
        return ""
    try:
        if isinstance(value, float) and math.isnan(value):
            return ""
    except Exception:
        pass
    s = str(value).strip()
    return "" if s.lower() == "nan" else s


def _is_drive_url(url: str) -> bool:
    if not url:
        return False
    low = url.lower()
    return "drive.google.com" in low or "docs.google.com" in low


def _extract_drive_file_id(url: str) -> Optional[str]:
    if not url:
        return None
    for pattern in [
        r"/file/d/([a-zA-Z0-9-_]+)",
        r"id=([a-zA-Z0-9-_]+)",
        r"drive\.google\.com/open\?id=([a-zA-Z0-9-_]+)",
    ]:
        m = re.search(pattern, url)
        if m:
            return m.group(1)
    return None


def _is_youtube_url(url: str) -> bool:
    if not url:
        return False
    low = url.lower()
    return "youtube.com" in low or "youtu.be" in low


def _parse_youtube_url(url: str) -> Tuple[Optional[str], Optional[int], Optional[int]]:
    if not url:
        return None, None, None
    video_id = None
    if "youtu.be/" in url:
        video_id = url.split("youtu.be/")[-1].split("?")[0].split("/")[0]
    elif "youtube.com" in url:
        if "/embed/" in url:
            video_id = url.split("/embed/")[-1].split("?")[0].split("/")[0]
        elif "v=" in url:
            parsed = urlparse(url)
            params = parse_qs(parsed.query)
            video_id = (params.get("v") or [None])[0]
    if not video_id:
        return None, None, None
    parsed = urlparse(url)
    params = parse_qs(parsed.query)
    start = end = None
    if "start" in params:
        try:
            start = int(params["start"][0])
        except (ValueError, TypeError):
            pass
    elif "t" in params:
        try:
            t_val = params["t"][0]
            start = int(t_val[:-1] if isinstance(t_val, str) and t_val.endswith("s") else t_val)
        except (ValueError, TypeError):
            pass
    if "end" in params:
        try:
            end = int(params["end"][0])
        except (ValueError, TypeError):
            pass
    return video_id, start, end


def _detect_asset_type(url: str) -> str:
    if not url:
        return "unknown"
    if _is_youtube_url(url):
        return "video"
    return "image"


def _ist_now() -> datetime:
    """Current time in IST (UTC+5:30)."""
    return datetime.now(timezone(timedelta(hours=5, minutes=30)))


def _filename_timestamp_ist() -> str:
    """Filesystem-safe IST timestamp suffix, aligned with Download Assets naming intent."""
    t = _ist_now()
    return t.strftime("%d-%m-%Y_%H-%M_IST")


# ---------------------------------------------------------------------------
# Parsing final_graphics_definition
# ---------------------------------------------------------------------------

# Reuse the battle-tested parser from review_and_revise
from agents.graphics_definition_v2.review_agent.review_and_revise import (
    parse_final_graphics_definition,
    parse_visual_steps,
)


def _flatten_visual_steps_from_text(text: str) -> List[Dict[str, Any]]:
    """
    Parse final_graphics_definition text into a flat list of visual steps.
    Each entry: {"segment_num", "step_index", "voiceover", "asset", "visual_id"}.
    """
    out: List[Dict[str, Any]] = []
    cleaned = _safe_str(text)
    if not cleaned:
        return out
    segments = parse_final_graphics_definition(cleaned)
    if not segments:
        return out
    for segment_num in sorted(segments.keys()):
        segment_text = segments[segment_num]
        steps = parse_visual_steps(segment_text)
        for idx, step in enumerate(steps, start=1):
            asset = _safe_str(step.get("asset", ""))
            voiceover = _safe_str(step.get("voiceover_part", ""))
            out.append({
                "segment_num": segment_num,
                "step_index": idx,
                "visual_id": f"S{segment_num}V{idx}",
                "voiceover": voiceover,
                "asset": asset,
            })
    return out


def collect_all_visual_steps(sheet) -> Tuple[List[Dict[str, Any]], str]:
    """
    Read every row of the Slide Chunks worksheet and return a flat list of
    visual steps plus the course name from Course info.
    """
    _, df = get_sheet_data_and_df(sheet, DEFAULT_WORKSHEET)
    if FINAL_GRAPHICS_COLUMN not in df.columns:
        raise ValueError(f"Column '{FINAL_GRAPHICS_COLUMN}' not found in worksheet.")

    _, course_info_df = get_sheet_data_and_df(sheet, "Course info")
    course_name = ""
    if not course_info_df.empty and "Course Name" in course_info_df.columns:
        course_name = _safe_str(course_info_df.loc[0, "Course Name"])

    all_steps: List[Dict[str, Any]] = []
    for row_idx, row in df.iterrows():
        raw = _safe_str(row.get(FINAL_GRAPHICS_COLUMN, ""))
        if not raw:
            continue
        slide_1based = int(row_idx) + 1
        steps = _flatten_visual_steps_from_text(raw)
        for s in steps:
            s["slide_index"] = slide_1based
            all_steps.append(s)
    return all_steps, course_name


# ---------------------------------------------------------------------------
# Edge TTS audio generation
# ---------------------------------------------------------------------------

async def _generate_tts_async(text: str, output_path: str, voice: str) -> str:
    communicate = edge_tts.Communicate(text, voice)
    await communicate.save(output_path)
    return output_path


def generate_tts(text: str, output_path: str, voice: str = DEFAULT_VOICE) -> str:
    loop = asyncio.get_event_loop()
    return loop.run_until_complete(_generate_tts_async(text, output_path, voice))


def generate_tts_batch(
    items: List[Tuple[str, str]], voice: str = DEFAULT_VOICE
) -> None:
    """Generate TTS for [(text, output_path), ...] in parallel."""
    async def _batch():
        tasks = [_generate_tts_async(text, path, voice) for text, path in items]
        await asyncio.gather(*tasks)
    loop = asyncio.get_event_loop()
    loop.run_until_complete(_batch())


def get_audio_duration_mutagen(path: str) -> float:
    """Get MP3 duration using mutagen (lightweight, no moviepy dependency)."""
    try:
        from mutagen.mp3 import MP3
        audio = MP3(path)
        return audio.info.length
    except Exception:
        pass
    try:
        from moviepy.editor import AudioFileClip
        c = AudioFileClip(path)
        d = c.duration
        c.close()
        return d
    except Exception:
        return 0.0


def _resolve_audio_duration(step: Dict[str, Any]) -> Tuple[str, float]:
    """
    Worker helper for parallel duration extraction.
    Returns (audio_path, duration_seconds) for mapping back to steps.
    """
    ap = step.get("_audio_path")
    if ap and os.path.exists(ap):
        return ap, get_audio_duration_mutagen(ap)
    return "", 2.5


# ---------------------------------------------------------------------------
# Asset download helpers
# ---------------------------------------------------------------------------

def _download_drive_image(drive: GoogleDrive, file_id: str, dest: str) -> Optional[str]:
    if not drive or not file_id:
        return None
    try:
        f = drive.CreateFile({"id": file_id})
        f.GetContentFile(dest)
        if os.path.exists(dest) and os.path.getsize(dest) > 0:
            return dest
    except Exception as e:
        print(f"  Drive download failed ({file_id}): {e}")
    return None


def _download_web_image(url: str, dest: str) -> Optional[str]:
    try:
        r = requests.get(url, timeout=30, stream=True)
        r.raise_for_status()
        with open(dest, "wb") as f:
            for chunk in r.iter_content(8192):
                f.write(chunk)
        if os.path.exists(dest) and os.path.getsize(dest) > 0:
            return dest
    except Exception as e:
        print(f"  Web image download failed ({url[:80]}): {e}")
    return None


def _download_youtube_clip(url: str, output_path: str) -> Optional[str]:
    """Download YouTube clip with optional start/end trimming."""
    video_id, start, end = _parse_youtube_url(url)
    if not video_id:
        return None
    youtube_url = f"https://www.youtube.com/watch?v={video_id}"
    tmp_dir = tempfile.gettempdir()
    base = os.path.join(tmp_dir, f"pptx_yt_{video_id}_{os.getpid()}")
    out_tmpl = base + ".%(ext)s"

    if start is not None and end is not None and end > start:
        cmd = [
            "yt-dlp",
            "-f", "bestvideo[height<=720][ext=mp4]+bestaudio[ext=m4a]/best[height<=720][ext=mp4]/best[height<=720]",
            "--extractor-args", "youtube:player_client=android",
            "--download-sections", f"*{start}-{end}",
            "--merge-output-format", "mp4",
            "-o", out_tmpl,
            "--no-playlist",
            youtube_url,
        ]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        if r.returncode == 0:
            for ext in ["mp4", "webm", "mkv"]:
                p = base + "." + ext
                if os.path.exists(p) and os.path.getsize(p) > 0:
                    if p != output_path:
                        shutil.move(p, output_path)
                    if os.path.exists(output_path):
                        return output_path
                    return p

    # Fallback: full download + trim
    full_tmpl = base + "_full.%(ext)s"
    cmd = [
        "yt-dlp",
        "-f", "bestvideo[height<=720][ext=mp4]+bestaudio[ext=m4a]/best[height<=720][ext=mp4]/best[height<=720]",
        "--extractor-args", "youtube:player_client=android",
        "--merge-output-format", "mp4",
        "-o", full_tmpl,
        "--no-playlist",
        youtube_url,
    ]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    if r.returncode != 0:
        return None
    actual_full = None
    for ext in ["mp4", "webm", "mkv"]:
        p = base + "_full." + ext
        if os.path.exists(p) and os.path.getsize(p) > 0:
            actual_full = p
            break
    if not actual_full:
        return None

    if start is not None and end is not None and end > start:
        duration = end - start
        trim_cmd = [
            "ffmpeg", "-y",
            "-ss", str(start), "-i", actual_full,
            "-t", str(duration),
            "-c", "copy", "-avoid_negative_ts", "make_zero",
            "-movflags", "+faststart",
            output_path,
        ]
        subprocess.run(trim_cmd, capture_output=True, text=True, timeout=120)
        try:
            os.remove(actual_full)
        except Exception:
            pass
        return output_path if os.path.exists(output_path) and os.path.getsize(output_path) > 0 else None

    shutil.move(actual_full, output_path)
    return output_path if os.path.exists(output_path) else None


def _trim_video_to_duration(video_path: str, max_duration: float, temp_dir: str) -> str:
    """Re-trim a downloaded video to max_duration seconds (narration length)."""
    if max_duration <= 0:
        return video_path
    try:
        from moviepy.editor import VideoFileClip
        clip = VideoFileClip(video_path)
        vid_dur = clip.duration
        clip.close()
        if vid_dur <= max_duration + 0.5:
            return video_path
    except Exception:
        return video_path

    trimmed = os.path.join(temp_dir, f"trimmed_{os.path.basename(video_path)}")
    cmd = [
        "ffmpeg", "-y",
        "-i", video_path,
        "-t", f"{max_duration:.2f}",
        "-c", "copy",
        "-avoid_negative_ts", "make_zero",
        "-movflags", "+faststart",
        trimmed,
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    if result.returncode == 0 and os.path.exists(trimmed) and os.path.getsize(trimmed) > 0:
        return trimmed
    return video_path


def _convert_video_to_gif(
    video_path: str,
    temp_dir: str,
    fps: int = 12,
    max_width: int = 960,
    target_duration_sec: Optional[float] = None,
) -> Optional[str]:
    """
    Convert an MP4 clip to animated GIF for reliable autoplay in PowerPoint.
    Returns GIF path or None if conversion fails.
    """
    if not video_path or not os.path.exists(video_path):
        return None
    gif_path = os.path.join(
        temp_dir,
        f"{os.path.splitext(os.path.basename(video_path))[0]}_autoplay.gif",
    )
    palette = os.path.join(
        temp_dir,
        f"{os.path.splitext(os.path.basename(video_path))[0]}_palette.png",
    )

    # Build filter chain. If narration is longer than source clip,
    # freeze the final frame so visual motion does not restart/loop.
    vf = f"fps={fps},scale={max_width}:-1:flags=lanczos"
    if target_duration_sec and target_duration_sec > 0:
        try:
            probe_cmd = [
                "ffprobe", "-v", "error",
                "-show_entries", "format=duration",
                "-of", "csv=p=0",
                video_path,
            ]
            probe = subprocess.run(probe_cmd, capture_output=True, text=True, timeout=30)
            clip_dur = float(probe.stdout.strip()) if probe.returncode == 0 and probe.stdout.strip() else None
            if clip_dur is not None and target_duration_sec > clip_dur + 0.05:
                freeze_sec = max(0.0, target_duration_sec - clip_dur)
                vf = f"{vf},tpad=stop_mode=clone:stop_duration={freeze_sec:.2f}"
        except Exception:
            pass

    # Palette generation improves GIF quality and reduces banding.
    palette_cmd = [
        "ffmpeg", "-y",
        "-i", video_path,
        "-vf", f"{vf},palettegen",
        palette,
    ]
    p1 = subprocess.run(palette_cmd, capture_output=True, text=True, timeout=180)
    if p1.returncode != 0 or not os.path.exists(palette):
        return None

    gif_cmd = [
        "ffmpeg", "-y",
        "-i", video_path,
        "-i", palette,
        "-lavfi", f"{vf}[x];[x][1:v]paletteuse",
        "-loop", "-1",
        gif_path,
    ]
    p2 = subprocess.run(gif_cmd, capture_output=True, text=True, timeout=240)
    try:
        if os.path.exists(palette):
            os.remove(palette)
    except Exception:
        pass
    if p2.returncode == 0 and os.path.exists(gif_path) and os.path.getsize(gif_path) > 0:
        return gif_path
    return None


# ---------------------------------------------------------------------------
# PPTX construction
# ---------------------------------------------------------------------------

def _image_extension(path: str) -> str:
    """Guess image extension; ensures we use a format pptx supports."""
    try:
        img = Image.open(path)
        fmt = (img.format or "").lower()
        img.close()
        if fmt in ("jpeg", "jpg"):
            return "jpg"
        if fmt == "png":
            return "png"
        if fmt == "gif":
            return "gif"
        if fmt == "webp":
            converted = path + ".png"
            img = Image.open(path).convert("RGB")
            img.save(converted, "PNG")
            img.close()
            return "png"
        return "png"
    except Exception:
        return "jpg"


def _convert_to_pptx_friendly(path: str) -> str:
    """Convert WEBP / unusual formats to PNG so python-pptx can embed them."""
    try:
        img = Image.open(path)
        fmt = (img.format or "").lower()
        if fmt in ("jpeg", "jpg", "png", "gif", "bmp", "tiff"):
            img.close()
            return path
        converted = path + ".png"
        img.convert("RGB").save(converted, "PNG")
        img.close()
        return converted
    except Exception:
        return path


def _add_image_slide(prs: Presentation, image_path: str, audio_path: Optional[str], duration_ms: int):
    """Add a slide with a centered/scaled image and optional narration audio."""
    slide_layout = prs.slide_layouts[6]  # blank layout
    slide = prs.slides.add_slide(slide_layout)

    friendly_path = _convert_to_pptx_friendly(image_path)

    try:
        img = Image.open(friendly_path)
        img_w_px, img_h_px = img.size
        img.close()
    except Exception:
        img_w_px, img_h_px = 1920, 1080

    # Convert pixel dimensions to EMU (assuming 96 DPI) and scale to fit slide
    img_w_emu = Inches(img_w_px / 96.0)
    img_h_emu = Inches(img_h_px / 96.0)
    scale = min(SLIDE_WIDTH_EMU / img_w_emu, SLIDE_HEIGHT_EMU / img_h_emu)
    pic_w = int(img_w_emu * scale)
    pic_h = int(img_h_emu * scale)
    left = (SLIDE_WIDTH_EMU - pic_w) // 2
    top = (SLIDE_HEIGHT_EMU - pic_h) // 2

    slide.shapes.add_picture(friendly_path, left, top, pic_w, pic_h)

    if audio_path and os.path.exists(audio_path):
        _embed_audio(slide, audio_path)

    _set_slide_auto_advance(slide, duration_ms)


def _add_video_slide(prs: Presentation, video_path: str, audio_path: Optional[str], duration_ms: int):
    """
    Add a slide for a video asset.
    Uses animated GIF for reliable autoplay in PowerPoint; falls back to embedded MP4.
    """
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    try:
        from moviepy.editor import VideoFileClip
        clip = VideoFileClip(video_path)
        vid_w_px, vid_h_px = clip.size
        clip.close()
    except Exception:
        vid_w_px, vid_h_px = 1280, 720

    vid_w_emu = Inches(vid_w_px / 96.0)
    vid_h_emu = Inches(vid_h_px / 96.0)
    scale = min(SLIDE_WIDTH_EMU / vid_w_emu, SLIDE_HEIGHT_EMU / vid_h_emu)
    pic_w = int(vid_w_emu * scale)
    pic_h = int(vid_h_emu * scale)
    left = (SLIDE_WIDTH_EMU - pic_w) // 2
    top = (SLIDE_HEIGHT_EMU - pic_h) // 2

    gif_path = _convert_video_to_gif(
        video_path,
        tempfile.gettempdir(),
        target_duration_sec=max(0.1, duration_ms / 1000.0),
    )
    if gif_path and os.path.exists(gif_path):
        slide.shapes.add_picture(gif_path, left, top, pic_w, pic_h)
    else:
        # Fallback to embedded movie if GIF conversion fails.
        poster_path = video_path + "_poster.jpg"
        try:
            from moviepy.editor import VideoFileClip
            clip = VideoFileClip(video_path)
            frame = clip.get_frame(0)
            clip.close()
            Image.fromarray(frame).save(poster_path, "JPEG")
        except Exception:
            img = Image.new("RGB", (vid_w_px, vid_h_px), (0, 0, 0))
            img.save(poster_path, "JPEG")
        slide.shapes.add_movie(
            video_path,
            left, top, pic_w, pic_h,
            poster_frame_image=poster_path,
            mime_type="video/mp4",
        )

    if audio_path and os.path.exists(audio_path):
        _embed_audio(slide, audio_path)

    _set_slide_auto_advance(slide, duration_ms)


def _add_placeholder_slide(prs: Presentation, error_text: str, audio_path: Optional[str], duration_ms: int):
    """Add a slide with error text when asset download fails."""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    txBox = slide.shapes.add_textbox(
        Inches(1), Inches(2.5), Inches(11.333), Inches(2.5)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = error_text
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(255, 80, 80)
    p.alignment = PP_ALIGN.CENTER

    if audio_path and os.path.exists(audio_path):
        _embed_audio(slide, audio_path)

    _set_slide_auto_advance(slide, duration_ms)


def _embed_audio(slide, audio_path: str):
    """Embed an MP3 audio file into a slide (hidden, auto-play)."""
    left = Inches(0)
    top = Inches(0)
    width = Emu(1)
    height = Emu(1)
    slide.shapes.add_movie(
        audio_path, left, top, width, height,
        mime_type="audio/mpeg",
    )
    _force_slide_media_autoplay(slide)


def _force_slide_media_autoplay(slide):
    """
    Force all media timing conditions on this slide to autoplay.
    python-pptx writes media start condition as delay='indefinite' (click-to-play).
    We rewrite it to delay='0' so narration starts automatically.
    """
    from pptx.oxml.ns import qn

    sld = slide._element
    timing = sld.find(qn("p:timing"))
    if timing is None:
        return

    # Media nodes are represented as p:video with p:cMediaNode.
    videos = timing.findall(".//" + qn("p:video"))
    for v in videos:
        for cond in v.findall(".//" + qn("p:cond")):
            cond.set("delay", "0")


def _set_slide_auto_advance(slide, duration_ms: int):
    """Set slide to auto-advance after duration_ms milliseconds."""
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    sld = slide._element
    transition = sld.find(qn("p:transition"))
    if transition is None:
        transition = OxmlElement("p:transition")
        # Keep a stable/valid order: place transition before timing if timing exists.
        timing = sld.find(qn("p:timing"))
        if timing is not None:
            sld.insert(list(sld).index(timing), transition)
        else:
            sld.append(transition)
    transition.set("advTm", str(max(int(duration_ms), 500)))
    transition.set("advClick", "0")


# ---------------------------------------------------------------------------
# Google Drive helpers
# ---------------------------------------------------------------------------

def _extract_sheet_id(url: str) -> Optional[str]:
    m = re.search(r"/d/([a-zA-Z0-9_-]+)", url)
    return m.group(1) if m else None


def _get_sheet_parent_folder_id(drive: GoogleDrive, sheet_id: str) -> Optional[str]:
    """Get the parent folder ID of a Google Sheet using the Drive API."""
    try:
        f = drive.CreateFile({"id": sheet_id})
        f.FetchMetadata(fields="parents")
        parents = f.get("parents", [])
        if parents:
            parent = parents[0]
            return parent.get("id") if isinstance(parent, dict) else parent
    except Exception as e:
        print(f"  Could not determine parent folder of sheet: {e}")
    return None


def _upload_file_to_drive(drive: GoogleDrive, parent_folder_id: str, local_path: str, filename: str) -> Optional[str]:
    """Upload a file to Drive and return its shareable URL"""
    max_attempts = 4
    backoff_seconds = [2, 5, 10]
    last_error = None

    for attempt in range(1, max_attempts + 1):
        try:
            f = drive.CreateFile({
                "title": filename,
                "parents": [{"id": parent_folder_id}],
            })
            f.SetContentFile(local_path)
            f.Upload()
            file_id = f["id"]
            return f"https://drive.google.com/file/d/{file_id}/view"
        except Exception as e:
            last_error = e
            print(f"  Upload attempt {attempt}/{max_attempts} failed: {e}")
            if attempt < max_attempts:
                sleep_s = backoff_seconds[min(attempt - 1, len(backoff_seconds) - 1)]
                time.sleep(sleep_s)

    print(f"  Upload failed after {max_attempts} attempts: {last_error}")
    return None


# ---------------------------------------------------------------------------
# Main orchestration
# ---------------------------------------------------------------------------

def generate_pptx_from_sheet(
    sheet_link: str,
    gc: gspread.Client,
    drive: GoogleDrive,
    voice: str = DEFAULT_VOICE,
    worksheet_name: str = DEFAULT_WORKSHEET,
    progress_callback=None,
) -> Optional[str]:
    """
    Generate a PPTX from final_graphics_definition and upload to the
    sheet's parent Drive folder.

    Returns the Drive URL of the uploaded PPTX, or None on failure.
    """

    def _progress(msg: str, fraction: Optional[float] = None):
        print(f"  [PPTX] {msg}")
        if progress_callback:
            try:
                progress_callback(msg, fraction)
            except TypeError:
                # Backward compatibility for older one-arg callbacks
                progress_callback(msg)

    _progress("Opening sheet...", 0.02)
    sheet_id = _extract_sheet_id(sheet_link)
    if not sheet_id:
        raise ValueError(f"Could not extract sheet ID from: {sheet_link}")
    sheet = gc.open_by_key(sheet_id)

    _progress("Reading visual steps from final_graphics_definition...", 0.08)
    all_steps, course_name = collect_all_visual_steps(sheet)
    if not all_steps:
        raise ValueError("No visual steps found in final_graphics_definition column.")
    _progress(f"Found {len(all_steps)} visual steps across sheet.", 0.12)

    safe_course_name = re.sub(r'[<>:"/\\|?*]', "_", course_name.strip()) if course_name else "Course"
    pptx_filename = (
        f"{safe_course_name}_Graphics_Definition_Slides_{_filename_timestamp_ist()}.pptx"
    )

    temp_dir = tempfile.mkdtemp(prefix="pptx_export_")
    audio_dir = os.path.join(temp_dir, "audio")
    asset_dir = os.path.join(temp_dir, "assets")
    os.makedirs(audio_dir, exist_ok=True)
    os.makedirs(asset_dir, exist_ok=True)

    try:
        # Phase 1: Generate all TTS audio in parallel
        _progress("Generating TTS narration audio...", 0.15)
        tts_items: List[Tuple[str, str]] = []
        for i, step in enumerate(all_steps):
            vo = step.get("voiceover", "").strip()
            if vo:
                audio_path = os.path.join(audio_dir, f"audio_{i}.mp3")
                tts_items.append((vo, audio_path))
                step["_audio_path"] = audio_path
            else:
                step["_audio_path"] = None
        if tts_items:
            generate_tts_batch(tts_items, voice)
        _progress(f"Generated {len(tts_items)} audio files.", 0.38)

        # Phase 2: Get audio durations in parallel
        _progress("Getting audio durations...", 0.42)
        duration_map: Dict[str, float] = {}
        with ThreadPoolExecutor(max_workers=_io_worker_count()) as executor:
            futures = [executor.submit(_resolve_audio_duration, step) for step in all_steps]
            done_count = 0
            for future in as_completed(futures):
                ap, dur = future.result()
                if ap:
                    duration_map[ap] = dur
                done_count += 1
                dur_fraction = 0.42 + (0.03 * (done_count / max(1, len(all_steps))))
                _progress(f"Reading audio durations ({done_count}/{len(all_steps)})...", dur_fraction)
        for step in all_steps:
            ap = step.get("_audio_path")
            step["_audio_duration"] = duration_map.get(ap, 2.5)

        # Phase 3: Download assets in parallel
        _progress(f"Downloading {len(all_steps)} assets...", 0.45)

        def _download_one(idx_step):
            idx, step = idx_step
            asset = step.get("asset", "")
            asset_type = _detect_asset_type(asset)
            narration_dur = step.get("_audio_duration", 2.5)

            if asset_type == "video":
                dest = os.path.join(asset_dir, f"video_{idx}.mp4")
                path = _download_youtube_clip(asset, dest)
                if path:
                    path = _trim_video_to_duration(path, narration_dur, asset_dir)
                return idx, path, "video"
            else:
                dest = os.path.join(asset_dir, f"image_{idx}.bin")
                if _is_drive_url(asset):
                    fid = _extract_drive_file_id(asset)
                    path = _download_drive_image(drive, fid, dest) if fid else None
                else:
                    path = _download_web_image(asset, dest)
                return idx, path, "image"

        asset_results: Dict[int, Tuple[Optional[str], str]] = {}
        with ThreadPoolExecutor(max_workers=_io_worker_count()) as executor:
            futures = {executor.submit(_download_one, (i, s)): i for i, s in enumerate(all_steps)}
            done_count = 0
            for future in as_completed(futures):
                idx, path, atype = future.result()
                asset_results[idx] = (path, atype)
                done_count += 1
                download_fraction = 0.45 + (0.20 * (done_count / max(1, len(all_steps))))
                _progress(f"Downloading assets ({done_count}/{len(all_steps)})...", download_fraction)
        _progress("All assets downloaded.", 0.65)

        # Phase 4: Build PPTX
        _progress("Building PowerPoint presentation...", 0.68)
        prs = Presentation()
        prs.slide_width = SLIDE_WIDTH_EMU
        prs.slide_height = SLIDE_HEIGHT_EMU

        for i, step in enumerate(all_steps):
            asset_path, asset_type = asset_results.get(i, (None, "image"))
            audio_path = step.get("_audio_path")
            duration_sec = step.get("_audio_duration", 2.5)
            duration_ms = int(duration_sec * 1000)

            slide_label = f"Slide {step.get('slide_index', '?')}, {step.get('visual_id', '?')}"

            if asset_path and os.path.exists(asset_path):
                if asset_type == "video":
                    try:
                        _add_video_slide(prs, asset_path, audio_path, duration_ms)
                        build_fraction = 0.68 + (0.24 * ((i + 1) / max(1, len(all_steps))))
                        _progress(f"Building slides ({i+1}/{len(all_steps)})...", build_fraction)
                    except Exception as e:
                        _add_placeholder_slide(
                            prs,
                            f"{slide_label}\nVideo embed failed: {e}\nAsset: {step.get('asset', '')}",
                            audio_path, duration_ms,
                        )
                        _progress(f"  [{i+1}/{len(all_steps)}] {slide_label} - video embed failed, placeholder added")
                else:
                    try:
                        _add_image_slide(prs, asset_path, audio_path, duration_ms)
                        build_fraction = 0.68 + (0.24 * ((i + 1) / max(1, len(all_steps))))
                        _progress(f"Building slides ({i+1}/{len(all_steps)})...", build_fraction)
                    except Exception as e:
                        _add_placeholder_slide(
                            prs,
                            f"{slide_label}\nImage embed failed: {e}\nAsset: {step.get('asset', '')}",
                            audio_path, duration_ms,
                        )
                        _progress(f"  [{i+1}/{len(all_steps)}] {slide_label} - image embed failed, placeholder added")
            else:
                _add_placeholder_slide(
                    prs,
                    f"{slide_label}\nAsset download failed.\nURL: {step.get('asset', 'N/A')}",
                    audio_path, duration_ms,
                )
                build_fraction = 0.68 + (0.24 * ((i + 1) / max(1, len(all_steps))))
                _progress(f"Building slides ({i+1}/{len(all_steps)})...", build_fraction)

        pptx_local = os.path.join(temp_dir, pptx_filename)
        prs.save(pptx_local)
        _progress("PPTX generated locally.", 0.94)

        # Phase 5: Upload to Drive parent folder
        _progress("Uploading PPTX to Google Drive...", 0.96)
        parent_folder_id = _get_sheet_parent_folder_id(drive, sheet_id)
        if not parent_folder_id:
            _progress("WARNING: Could not find parent folder, uploading to root.")
            parent_folder_id = "root"

        drive_url = _upload_file_to_drive(drive, parent_folder_id, pptx_local, pptx_filename)
        if drive_url:
            _progress("Upload complete.", 1.0)
        else:
            _progress(
                f"Upload failed after retries. Local PPTX was generated at: {pptx_local}"
            )
        return drive_url

    finally:
        try:
            shutil.rmtree(temp_dir, ignore_errors=True)
        except Exception:
            pass


# ---------------------------------------------------------------------------
# Streamlit UI
# ---------------------------------------------------------------------------

def main():
    import streamlit as st

    try:
        st.set_page_config(page_title="PPTX Exporter", page_icon="", layout="wide")
    except Exception:
        pass

    st.title("PPTX Exporter — Graphics Definition V2")
    st.caption(
        "Generate a PowerPoint presentation from final_graphics_definition. "
        "One visual per slide with Edge TTS narration."
    )

    if "gc" not in st.session_state:
        st.error("Google Sheets client not found. Please log in first.")
        return
    gc = st.session_state["gc"]
    drive = st.session_state.get("drive")
    if not drive:
        st.error("Google Drive instance not found. Please log in first.")
        return

    sheet_link = st.text_input("Google Sheet link", value=st.session_state.get("sheet_link", ""))
    if st.button("Generate PPTX", type="primary"):
        if not sheet_link:
            st.error("Please enter a sheet link.")
            return
        status_area = st.empty()
        smart_progress = SmartProgressBar(total_tasks=100, description="PPTX Export", save_interval=0)

        def _cb(msg: str, fraction: Optional[float] = None):
            if fraction is not None:
                smart_progress.update_progress(
                    fraction,
                    custom_text=f"PPTX Export: {int(max(0.0, min(1.0, fraction)) * 100)}% - {msg}",
                )
            status_area.caption(msg)

        try:
            with st.spinner("Generating PPTX..."):
                url = generate_pptx_from_sheet(
                    sheet_link=sheet_link,
                    gc=gc,
                    drive=drive,
                    voice=DEFAULT_VOICE,
                    progress_callback=_cb,
                )
            smart_progress.update_progress(1.0, custom_text="PPTX Export: 100% - Done")
            if url:
                st.success(f"PPTX uploaded to Drive: [Open]({url})")
            else:
                st.error("PPTX generation completed but upload failed. Check logs.")
        except Exception as e:
            st.error(f"Error: {e}")
            import traceback
            st.code(traceback.format_exc())

main()
