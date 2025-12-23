
import os
import re
import tempfile
from typing import List
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import ns
from docx.oxml.shared import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def parse_quality_summary_sections(text: str) -> List[dict]:
    """Parse the LLM quality summary into structured sections."""
    sections = []
    if text:
        pattern = re.compile(
            r"\d+\.\s+(?P<header>[^\n]+)\nComments Grouped:\s*(?P<comments>.*?)(?:\nActionable Summary:\s*(?P<action>.*?))?(?=\n\d+\.\s+|\Z)",
            re.S
        )
        for idx, match in enumerate(pattern.finditer(text.strip()), 1):
            header = match.group("header").strip()
            topics_match = re.search(r"\[(.*?)\]", header)
            topics = topics_match.group(1).strip() if topics_match else ""
            category = re.sub(r"\[.*?\]", "", header).strip() or f"Category {idx}"
            comments_block = (match.group("comments") or "").strip()
            comments = [ln.strip() for ln in comments_block.splitlines() if ln.strip()]
            actionable = (match.group("action") or "").strip() or "No actionable summary provided."
            sections.append({
                "category": category,
                "topics": topics,
                "comments": comments,
                "actionable": actionable
            })
    return sections


def create_course_review_doc(result_dict, spreadsheet_url):
    """
    Build a .docx Course Review Report with refined formatting:
    - All text black, no italics
    - Dynamic reporter name
    - Correct font sizes and headings
    - Hyperlinked comment counts and Data Source
    """

    def safe_str(value):
        if isinstance(value, str):
            return value.strip()
        if value is None:
            return ""
        return str(value).strip()

    def sanitize_filename(name):
        return re.sub(r"[\\/*?\"<>|:]", "_", name)

    # --- universal style helper ---
    def set_black_bold(run, size=None, bold=True):
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.italic = False
        run.bold = bold
        if size:
            run.font.size = Pt(size)

    def add_hyperlink(paragraph, text, url=None, anchor=None):
        """Add a clickable hyperlink to a paragraph (external or internal)."""
        if not url and not anchor:
            raise ValueError("Either a URL or an anchor must be provided for hyperlinks.")

        hyperlink = OxmlElement("w:hyperlink")
        if anchor:
            hyperlink.set(ns.qn("w:anchor"), anchor)
        else:
            part = paragraph.part
            r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
            hyperlink.set(ns.qn("r:id"), r_id)

        new_run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")

        # Blue underline style for links only
        u = OxmlElement("w:u")
        u.set(ns.qn("w:val"), "single")
        rPr.append(u)
        color = OxmlElement("w:color")
        color.set(ns.qn("w:val"), "0000FF")
        rPr.append(color)
        new_run.append(rPr)

        text_elem = OxmlElement("w:t")
        text_elem.text = text
        new_run.append(text_elem)
        hyperlink.append(new_run)
        paragraph._p.append(hyperlink)
        return hyperlink

    bookmark_counter = {"value": 0}

    def add_bookmark(paragraph, name: str):
        """Create a bookmark so table hyperlinks can jump to comment sections."""
        bookmark_counter["value"] += 1
        bookmark_id = str(bookmark_counter["value"])
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(ns.qn("w:id"), bookmark_id)
        bookmark_start.set(ns.qn("w:name"), name)
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(ns.qn("w:id"), bookmark_id)
        paragraph._p.insert(0, bookmark_start)
        paragraph._p.append(bookmark_end)

    # ---- Extract data ----
    course_name = safe_str(result_dict.get("Course Name", "Unnamed Course"))
    stage = safe_str(result_dict.get("Stage", "Unknown Stage"))
    creator = safe_str(result_dict.get("Creator Name", "N/A"))
    reviewer = safe_str(result_dict.get("Reviewer Name", "N/A"))
    quality_summary = safe_str(result_dict.get("Quality Summary"))
    flagged_basic = safe_str(result_dict.get("Flagged Items - Basic"))
    flagged_critical = safe_str(result_dict.get("Flagged Items - Critical"))
    compliance_basic = safe_str(result_dict.get("Compliance Summary - Basic"))
    compliance_critical = safe_str(result_dict.get("Compliance Summary - Critical"))
    comments_text = safe_str(result_dict.get("Comments"))
    data_source = f"Data Source: Checklist (v1.05): {course_name}"

    if not (compliance_basic or compliance_critical or quality_summary):
        print(f"Skipping report for {creator} — all summaries are empty.")
        return None, None

    temp_dir = tempfile.mkdtemp(prefix="course_review_")
    output_filename = sanitize_filename(f"Course Review Report - {creator} ({stage}).docx")
    output_path = os.path.join(temp_dir, output_filename)

    # ---- Helpers ----
    def style_table(table):
        table.style = "Table Grid"
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                if p.runs:
                    set_black_bold(p.runs[0], bold=True)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    def parse_rephrased_block(text):
        """
        Parse rephrased items AND merge rows:
        Same Category + Same Item → Combine all topics.
        """
        grouped = {}  # key: (category, item) → set(topics)

        current_category = None
        current_topic = None

        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue

            # Detect category + topics line
            cat_match = re.match(r"^(.*?)\s*\[(Topic[^\]]+)\]$", line)
            if cat_match:
                current_category = cat_match.group(1).strip()
                topics = [t.strip() for t in cat_match.group(2).split(",")]
                current_topic = topics
                continue

            # This is a flagged item line
            if current_category and current_topic:
                key = (current_category, line)
                grouped.setdefault(key, set()).update(current_topic)

        # Convert to list of rows
        items = []
        for (category, item), topics in grouped.items():
            items.append({
                "Category": category,
                "Flagged Item": item,
                "Relevant Topic": ", ".join(sorted(topics, key=lambda x: int(x.split()[-1])))
            })

        return items


    basic_items = parse_rephrased_block(flagged_basic)
    critical_items = parse_rephrased_block(flagged_critical)

    # ---- Create Document ----
    doc = Document()

    # --- Title ---
    title = doc.add_heading("Review Report", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        set_black_bold(run, size=26, bold=True)

    # --- Subtitle (Reporter / Creator) ---
    subtitle = doc.add_paragraph(creator)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        set_black_bold(run, size=15, bold=False)
    doc.add_paragraph("")

    # --- Course Details ---
    for label, value in [("Course", course_name), ("Stage", stage), ("Moderator", reviewer)]:
        p = doc.add_paragraph()
        p.style = doc.styles["Heading 3"]
        run_label = p.add_run(f"{label}: ")
        set_black_bold(run_label, size=18, bold=True)
        run_value = p.add_run(value)
        set_black_bold(run_value, size=18, bold=False)

    doc.add_paragraph("")

    # --- Compliance Summary ---
    h = doc.add_heading("Compliance Summary", level=4)
    for run in h.runs:
        set_black_bold(run, size=18, bold=True)

    if basic_items:
        h = doc.add_heading("Basic Items Flagged:", level=5)
        for run in h.runs:
            set_black_bold(run, size=16, bold=True)
        table_basic = doc.add_table(rows=1, cols=3)
        hdr = table_basic.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Category", "Flagged Item", "Relevant Topic"
        style_table(table_basic)
        for item in basic_items:
            row = table_basic.add_row().cells
            row[0].text = item["Category"]
            row[1].text = item["Flagged Item"]
            row[2].text = item["Relevant Topic"]
        doc.add_paragraph("")

    if critical_items:
        h = doc.add_heading("Critical Items Flagged:", level=5)
        for run in h.runs:
            set_black_bold(run, size=16, bold=True)
        table_critical = doc.add_table(rows=1, cols=3)
        hdr = table_critical.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Category", "Flagged Item", "Relevant Topic"
        style_table(table_critical)
        for item in critical_items:
            row = table_critical.add_row().cells
            row[0].text = item["Category"]
            row[1].text = item["Flagged Item"]
            row[2].text = item["Relevant Topic"]
        doc.add_paragraph("")

    # --- Quality Summary ---
    if quality_summary:
        h = doc.add_heading("Quality Summary", level=4)
        for run in h.runs:
            set_black_bold(run, size=18, bold=True)

        sections = parse_quality_summary_sections(quality_summary)
        if not sections:
            doc.add_paragraph(quality_summary)
            doc.add_paragraph("")
        else:
            table_quality = doc.add_table(rows=1, cols=4)
            hdr = table_quality.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = (
                "Category", "Comments", "Improvements Needed", "Relevant Topic")
            style_table(table_quality)

            comment_targets = []

            for idx, section in enumerate(sections, 1):
                category = section.get("category") or f"Category {idx}"
                topics_text = section.get("topics", "")
                improvements_content = section.get("actionable", "")
                comment_list = section.get("comments") or []

                row = table_quality.add_row().cells
                row[0].text = category

                link_text = f"View {len(comment_list)} comment(s)"

                if comment_list:
                    safe_name = re.sub(r"[^A-Za-z0-9]+", "_", category).strip("_")
                    if not safe_name or not safe_name[0].isalpha():
                        safe_name = f"category_{idx}"
                    bookmark_name = f"{safe_name}_{idx}"
                    row[1].text = ""
                    p_link = row[1].paragraphs[0]
                    add_hyperlink(p_link, link_text, anchor=bookmark_name)
                    comment_targets.append({
                        "bookmark": bookmark_name,
                        "category": category,
                        "comments": comment_list
                    })
                else:
                    row[1].text = "No comments provided."

                row[2].text = improvements_content
                row[3].text = topics_text

            doc.add_paragraph("")

            if comment_targets:
                h = doc.add_heading("Comments", level=5)
                for run in h.runs:
                    set_black_bold(run, size=16, bold=True)

                for target in comment_targets:
                    cat_para = doc.add_paragraph(target["category"])
                    cat_para.style = doc.styles["Heading 6"]
                    add_bookmark(cat_para, target["bookmark"])
                    for run in cat_para.runs:
                        set_black_bold(run, size=14, bold=True)

                    for i, comment_text in enumerate(target["comments"], 1):
                        para = doc.add_paragraph(f"{i}. {comment_text}")
                        for run in para.runs:
                            set_black_bold(run, bold=False)

                doc.add_paragraph("")

    # --- Data Source hyperlink ---
    p = doc.add_paragraph()
    add_hyperlink(p, data_source, spreadsheet_url)

    # ---- Save ----
    doc.save(output_path)
    print(f"Course Review Report created: {output_path}")

    return output_path, {"temp_dir": temp_dir, "drive_name": f"Review Report - {creator} ({stage})"}
